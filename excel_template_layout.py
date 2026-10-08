"""Read native template shapes and align source text/photos without inventing a grid."""
import copy,json,math,posixpath,zipfile
from pathlib import Path
from lxml import etree as ET
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter,range_boundaries
from openpyxl.cell.cell import MergedCell
from auto_guide import save

S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
D='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
X='http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
NS={'x':X,'a':A}

def column_pixels(ws,column):
    key=get_column_letter(column)
    dimension=ws.column_dimensions.get(key)
    width=dimension.width if dimension else (ws.sheet_format.defaultColWidth or 8.43)
    return width*7+5

def drawing_part(z,sheet):
    node=ET.fromstring(z.read(sheet)).find('{'+S+'}drawing')
    if node is None:return None
    rel=posixpath.join(posixpath.dirname(sheet),'_rels',posixpath.basename(sheet)+'.rels')
    target=next(r.get('Target') for r in ET.fromstring(z.read(rel)) if r.get('Id')==node.get('{'+D+'}id'))
    return target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(sheet),target))

def template_shapes(template,sheet_index):
    with zipfile.ZipFile(template) as z:
        part=drawing_part(z,f'xl/worksheets/sheet{sheet_index+1}.xml')
        root=ET.fromstring(z.read(part)) if part else []
    shapes={};catalog=[]
    for index,anchor in enumerate(root):
        text=''.join(anchor.xpath('.//a:t/text()',namespaces=NS))
        if not text:continue
        f=anchor.find('x:from',NS);t=anchor.find('x:to',NS)
        if f is None:continue
        key=f'SHAPE_{index:03d}';shapes[key]=anchor
        catalog.append({'id':key,'text':text,'col':int(f.find('x:col',NS).text),'row':int(f.find('x:row',NS).text)+1,
                        'end_col':int(t.find('x:col',NS).text) if t is not None else 0})
    return shapes,catalog

def inject_textboxes(template,destination,placements):
    """Copy actual native textbox formatting; replace text and align only anchors."""
    with zipfile.ZipFile(destination) as z:parts={n:z.read(n) for n in z.namelist()}
    for index,items in placements.items():
        shapes,_=template_shapes(template,index)
        with zipfile.ZipFile(destination) as z:part=drawing_part(z,f'xl/worksheets/sheet{index+1}.xml')
        if not part:raise ValueError('Missing body drawing for native textboxes')
        sheet=ET.fromstring(parts[f'xl/worksheets/sheet{index+1}.xml'])
        fmt=sheet.find('{'+S+'}sheetFormatPr')
        default_width=float(fmt.get('defaultColWidth','8.43')) if fmt is not None else 8.43
        default_height=float(fmt.get('defaultRowHeight','15')) if fmt is not None else 15
        widths={};heights={}
        cols=sheet.find('{'+S+'}cols')
        if cols is not None:
            for col in cols:
                for c in range(int(col.get('min')),int(col.get('max'))+1):widths[c]=float(col.get('width',default_width))*7+5
        for row in sheet.findall('{'+S+'}sheetData/{'+S+'}row'):
            heights[int(row.get('r'))]=float(row.get('ht',default_height))*96/72
        def point(col,row):
            return (round(sum(widths.get(c,default_width*7+5) for c in range(1,col+1))*9525),
                    round(sum(heights.get(r,default_height*96/72) for r in range(1,row))*9525))
        root=ET.fromstring(parts[part]);next_id=max([0]+[int(n.get('id')) for n in root.xpath('.//x:cNvPr',namespaces=NS)])+1
        for item in items:
            anchor=copy.deepcopy(shapes[item['shape_id']]);shape=anchor.find('x:sp',NS)
            if shape is None:raise ValueError('Template text must be a native shape')
            textbody=shape.find('x:txBody',NS);paragraph=textbody.find('a:p',NS)
            run=paragraph.find('a:r',NS) if paragraph is not None else None
            # Preserve native textbox, font, paragraph and run properties, not stale words.
            for p in list(textbody.findall('a:p',NS)):textbody.remove(p)
            for line in item['text'].split('\n'):
                p=ET.SubElement(textbody,'{'+A+'}p')
                if paragraph is not None and paragraph.find('a:pPr',NS) is not None:p.append(copy.deepcopy(paragraph.find('a:pPr',NS)))
                r=ET.SubElement(p,'{'+A+'}r')
                if run is not None and run.find('a:rPr',NS) is not None:r.append(copy.deepcopy(run.find('a:rPr',NS)))
                ET.SubElement(r,'{'+A+'}t').text=line
            transform=shape.find('x:spPr/a:xfrm',NS)
            if transform is not None:
                left,top=point(item['col'],item['row']);right,bottom=point(item['end_col'],item['end_row'])
                off=transform.find('a:off',NS);ext=transform.find('a:ext',NS)
                if off is not None:off.set('x',str(left));off.set('y',str(top))
                if ext is not None:ext.set('cx',str(right-left));ext.set('cy',str(bottom-top))
            # Shared anchors put both lanes on identical horizontal rows.
            for tag in ('from','to','ext','clientData'):
                for n in list(anchor.findall('x:'+tag,NS)):anchor.remove(n)
            for tag,col,row in [('from',item['col'],item['row']-1),('to',item['end_col'],item['end_row']-1)]:
                marker=ET.SubElement(anchor,'{'+X+'}'+tag)
                for k,v in [('col',col),('colOff',0),('row',row),('rowOff',0)]:ET.SubElement(marker,'{'+X+'}'+k).text=str(v)
            # Shape must follow from/to in SpreadsheetDrawing order.
            anchor.remove(shape);anchor.append(shape);ET.SubElement(anchor,'{'+X+'}clientData')
            for n in anchor.xpath('.//x:cNvPr',namespaces=NS):n.set('id',str(next_id));next_id+=1
            anchor.set('editAs','oneCell');root.append(anchor)
        parts[part]=ET.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    temporary=Path(destination).with_suffix('.aligned.tmp')
    with zipfile.ZipFile(temporary,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    temporary.replace(destination)

def native_layout(job,source,plan,pages,ask_helpers):
    obj,arr,string,photo_url=ask_helpers
    shapes,catalog=template_shapes(job.template,2)
    if not catalog:raise ValueError('Mẫu thiếu hộp chữ bố cục; cần mẫu có phần mô tả gốc.')
    cells=[{'cell':c.coordinate,'text':str(c.value)} for row in source for c in row if c.value is not None and c.row>4]
    shape_ids=[c['id'] for c in catalog];step_ids=[i['source_id'] for i in plan['items'] if i['kind']=='step']
    table_schema=obj({'range':string(30),'header_end_row':{'type':'integer','minimum':5,'maximum':source.max_row},'position':{'type':'string','enum':['top','side','bottom']},'keep':{'type':'boolean'}})
    lane_schema=obj({'heading_shape_id':{'type':'string','enum':shape_ids},'caption_shape_id':{'type':'string','enum':shape_ids}})
    schema=obj({'title_shape_ids':arr({'type':'string','enum':shape_ids},4),'tables':arr(table_schema,6),
                'lanes':arr(lane_schema,4,1),'assignment':obj({ident:{'type':'string','enum':['GENERAL']+shape_ids} for ident in step_ids})})
    prompt='''Read the NATIVE Excel layout, including floating textboxes. Preserve its table at the top, original left/right heading labels and original captions ABOVE photos in separate left/right lanes below the table. Do not make a new image/description column next to the wire table. Identify template tables by their actual range and last header row. Top=cut/measurement table; side=label/pin table; bottom=materials/BOM. Keep top/bottom headers; keep side only if input has supported matching data. Table ranges cover headers and their sample data, inclusive. Left/right lane heading_shape_id must point to the corresponding existing heading textbox. Caption prototype must be a full instruction textbox from that same lane, not a dimension or heading. Map source group A/B to left/right using the PDF, with source group label preserved separately; For assignment map each operation ID to exactly one lane heading_shape_id, or GENERAL for ungrouped preparation. Connector-side operations go to the connector side shown in PDF, not simply the first lane; the breakout/free-end operations go to the opposite side. Never map by source ordering alone. General steps without a group go above the two lanes. Retain native static title shape IDs only (e.g. table title), not sample instruction or dimensional labels. Return structural mapping only; code will use source text verbatim.\nTEMPLATE CELLS:\n'''+json.dumps(cells,ensure_ascii=False)+'\nNATIVE SHAPES:\n'+json.dumps(catalog,ensure_ascii=False)+'\nSOURCE:\n'+json.dumps([{'id':i['source_id'],'kind':i['kind'],'group':i.get('group_text',''),'text':i['text']} for i in plan['items']],ensure_ascii=False)
    def validate(d):
        heading_ids=[l['heading_shape_id'] for l in d['lanes']]
        if len(set(heading_ids))!=len(heading_ids):raise ValueError('Native lane headings must be unique')
        if any(v!='GENERAL' and v not in heading_ids for v in d['assignment'].values()):raise ValueError('Every operation must map to a real lane')
        for t in d['tables']:
            c1,r1,c2,r2=range_boundaries(t['range'])
            if not (1<=c1<=c2<=source.max_column and 5<=r1<=t['header_end_row']<r2<=source.max_row):raise ValueError('Invalid native table range')
        if len(set(d['title_shape_ids']))!=len(d['title_shape_ids']):raise ValueError('Duplicate title shape')
    content=[{'type':'text','text':prompt}]
    if pages:content.append({'type':'image_url','image_url':{'url':photo_url(pages[0]['path'],(1600,1200))}})
    layout=job.ask('native_floating_layout_v2',content,2500,validate,schema)
    for table in layout['tables']:
        # Include full native merged headers, even when the model stops at their
        # top-left cell. This changes geometry only, never source content.
        c1,r1,c2,r2=range_boundaries(table['range'])
        for m in source.merged_cells.ranges:
            if r1<=m.min_row and m.max_row<=r2 and c1<=m.min_col<=c2:c2=max(c2,m.max_col)
        table['range']=f'{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}'
    layout['general_source_ids']=[s for s in step_ids if layout['assignment'][s]=='GENERAL']
    for lane in layout['lanes']:lane['source_ids']=[s for s in step_ids if layout['assignment'][s]==lane['heading_shape_id']]
    for ordinal,table in enumerate(layout['tables']):
        c1,r1,c2,r2=range_boundaries(table['range'])
        headers=[c for c in cells if r1<=source[c['cell']].row<=table['header_end_row'] and c1<=source[c['cell']].column<=c2]
        if table['position']=='top':
            fields=['item_number','wire_part_number','wire_color','wire_cut_inches','wire_cut_mm','strip_tool','strip_left_inches','strip_right_inches','strip_left_mm','strip_right_mm']
            columns=list(range(c1,c2+1))
            if len(columns)!=len(fields):raise ValueError('Native wire headers need ten columns')
        elif table['position']=='side':
            fields=['wire_color','left_pin_or_label','right_printed_label'];columns=list(range(c1,c2+1))
            if len(columns)!=len(fields):raise ValueError('Native label table needs three columns')
        else:
            fields=['step','material_part_number','material_description','quantity','item_number','instruction','note']
            columns=sorted({source[c['cell']].column for c in headers if source[c['cell']].row==r1})
            if len(columns)!=len(fields):raise ValueError('Native materials table needs seven headers')
        body_rows=list(range(table['header_end_row']+1,r2+1))
        schema=obj({str(row):obj({field:string(300 if field in ('instruction','note') else 80) for field in fields}) for row in body_rows})
        relevant=[{'id':i['source_id'],'text':i['text']} for i in plan['items'] if i['kind'] in (('step','group') if table['position'] in ('top','side') else ('note',))]
        prompt='Extract typed records for this template table from SOURCE only. Return exactly the requested row keys and field keys. Unused rows and unsupported fields must be empty strings. Every nonempty value must be copied literally from source, not paraphrased. Never invent numbering or quantities. A complete assembly identifier is NOT wire_part_number. Source group headings are NOT wire colors or part numbers. Combine cut length, strip tool and left/right stripping sizes for the same unspecified wire in ONE row; group A/B is not a separate wire part. Strip dimensions are NOT printed labels or pin mappings. A tool is NOT a material. A source note about material replacement may supply the new material_part_number and an explicit item_number, but not invented description/quantity/step. Keep full source material notes only in note/instruction fields. Do not copy template sample data. All sample body values have been discarded.\nFIELDS: '+json.dumps(fields)+'\nNATIVE HEADERS: '+json.dumps(headers,ensure_ascii=False)+'\nLANE MEANINGS: '+json.dumps([{'heading':next(c['text'] for c in catalog if c['id']==l['heading_shape_id']),'source_ids':l['source_ids']} for l in layout['lanes']],ensure_ascii=False)+'\nSOURCE: '+json.dumps(relevant,ensure_ascii=False)
        records=job.ask('native_typed_table_v2_'+str(ordinal),prompt,3000,lambda d:None,schema)
        source_text=' '.join(i['text'] for i in relevant).casefold()
        metadata_values={i['text'].casefold() for i in plan['items'] if i['kind']=='metadata'}
        table['cells']=[];table['rejected_unquoted_cells']=[]
        for row,record in records.items():
            for column,field in zip(columns,fields):
                value=record[field]
                if not value:continue
                address=f'{get_column_letter(column)}{row}'
                if value.casefold() not in source_text or (field=='wire_part_number' and value.casefold() in metadata_values):
                    table['rejected_unquoted_cells'].append(address);continue
                # Only writable native cells are eligible; merged continuation cells
                # cannot receive fabricated copies of a value from another row.
                if not isinstance(source[address],MergedCell):table['cells'].append({'cell':address,'value':value})
    save(job.out/'native-layout.json',layout);return layout,catalog

def build_native_sheet(job,source,target,plan,pages,helpers):
    obj,arr,string,photo_url,copy_style,image_at=helpers
    layout,catalog=native_layout(job,source,plan,pages,(obj,arr,string,photo_url));shape_map={c['id']:c for c in catalog}
    placements=[];items={i['source_id']:i for i in plan['items']}
    def table(t,start=None):
        c1,r1,c2,r2=range_boundaries(t['range']);offset=(start-r1) if start else 0
        for row in source.iter_rows(min_row=r1,max_row=r2,min_col=c1,max_col=c2):
            for original in row:
                dest=target.cell(original.row+offset,original.column);copy_style(original,dest)
                if original.row<=t['header_end_row'] and original.value is not None:dest.value=original.value
        for m in source.merged_cells.ranges:
            if c1<=m.min_col and m.max_col<=c2 and r1<=m.min_row and m.max_row<=r2:
                target.merge_cells(start_row=m.min_row+offset,end_row=m.max_row+offset,start_column=m.min_col,end_column=m.max_col)
        for value in t['cells']:
            original=source[value['cell']];dest=target.cell(original.row+offset,original.column);dest.value=value['value'];dest.data_type='s'
            dest.alignment=Alignment(horizontal=original.alignment.horizontal,vertical='top',wrap_text=True)
            merged=next((m for m in target.merged_cells.ranges if dest.coordinate in m),None)
            last_col=merged.max_col if merged else dest.column
            pixels=sum(column_pixels(target,c) for c in range(dest.column,last_col+1))
            lines=sum(max(1,math.ceil(len(line)/max(8,pixels/7))) for line in value['value'].splitlines())
            target.row_dimensions[dest.row].height=max(target.row_dimensions[dest.row].height or 15,lines*15)
        return r2+offset+2
    row=5
    for t in layout['tables']:
        if t['keep'] and t['position']=='top':row=max(row,table(t))
    for t in layout['tables']:
        if t['keep'] and t['position']=='side' and t['cells']:row=max(row,table(t))
    for ident in layout['title_shape_ids']:
        shape=shape_map[ident]
        placements.append({'shape_id':ident,'text':shape['text'],'col':shape['col'],'end_col':max(shape['end_col'],shape['col']+3),'row':shape['row'],'end_row':shape['row']+2})
    lanes=sorted(layout['lanes'],key=lambda l:shape_map[l['heading_shape_id']]['col'])
    # Use the original image-lane bounds: broad left/right regions, never the wire columns.
    split=[shape_map[l['heading_shape_id']]['col'] for l in lanes]
    left_edges=[max(0,c-2) for c in split]
    if len(lanes)==2:left_edges=[0,max(1,split[1]-2)]
    spans=[(left_edges[i],(left_edges[i+1]-1 if i+1<len(lanes) else max(source.max_column,shape_map[lanes[i]['caption_shape_id']]['end_col']))) for i in range(len(lanes))]
    # Recover the image regions from the actual template, including columns beyond
    # the last populated cell. Those bounds are lost when only cells are inspected.
    boundaries=[(split[i]+split[i+1])/2 for i in range(len(split)-1)]
    for i in range(len(lanes)):
        pictures=[im.anchor for im in source._images if hasattr(im.anchor,'to') and im.anchor._from.row>=4
                  and (i==0 or im.anchor._from.col>=boundaries[i-1])
                  and (i==len(lanes)-1 or im.anchor._from.col<boundaries[i])
                  and im.anchor.to.col-im.anchor._from.col>=3]
        if pictures:spans[i]=(min(p._from.col for p in pictures),max(p.to.col+1 for p in pictures))
    for r in range(5,220):target.row_dimensions[r].height=15
    def width(c1,c2):return sum(column_pixels(target,c+1) for c in range(c1,c2))
    def add_step(ident,lane,start):
        item=items[ident];c1,c2=lane['bounds'];span=width(c1,c2)
        # Long source wording gets more room; no edited/shortened descriptions.
        text_rows=max(3,math.ceil(len(item['text'])/max(25,span/7))+1)
        placements.append({'shape_id':lane['caption_shape_id'],'text':item['text'],'source_id':ident,'col':c1,'end_col':c2,'row':start,'end_row':start+text_rows})
        if not item['image_ids']:raise ValueError('Every operation requires a Qwen-selected photo')
        image_at(target,job.image_map[item['image_ids'][0]]['path'],c1,start+text_rows,span-20,160)
        return text_rows+10
    base_lane=dict(lanes[0],bounds=spans[0])
    for ident in layout['general_source_ids']:row+=add_step(ident,base_lane,row)
    # Keep the two native headings at the same horizontal level, not different rows.
    row=max(row, min(shape_map[l['heading_shape_id']]['row'] for l in lanes))
    for i,lane in enumerate(lanes):
        c1,c2=spans[i];shape=shape_map[lane['heading_shape_id']]
        placements.append({'shape_id':lane['heading_shape_id'],'text':shape['text'],'col':c1,'end_col':c2,'row':row,'end_row':row+2})
        group=next((items[s].get('group_text','') for s in lane['source_ids'] if items[s].get('group_text')), '')
        if group:placements.append({'shape_id':lane['caption_shape_id'],'text':group,'col':c1,'end_col':c2,'row':row+2,'end_row':row+4})
    row+=5
    # Both caption and image tops share rows even for captions of unequal length.
    for ordinal in range(max(len(l['source_ids']) for l in lanes)):
        active=[(dict(l,bounds=spans[i]),l['source_ids'][ordinal]) for i,l in enumerate(lanes) if ordinal<len(l['source_ids'])]
        caption_rows=max(max(3,math.ceil(len(items[s]['text'])/max(25,width(*l['bounds'])/7))+1) for l,s in active)
        for lane,ident in active:
            item=items[ident];c1,c2=lane['bounds'];span=width(c1,c2)
            placements.append({'shape_id':lane['caption_shape_id'],'text':item['text'],'source_id':ident,'col':c1,'end_col':c2,'row':row,'end_row':row+caption_rows})
            if not item['image_ids']:raise ValueError('Every operation requires a photo')
            image_at(target,job.image_map[item['image_ids'][0]]['path'],c1,row+caption_rows,span-20,160)
        row+=caption_rows+10
    for t in layout['tables']:
        if t['keep'] and t['position']=='bottom':row=table(t,row+2)
    rendered_table_values=[c['value'] for t in layout['tables'] if t['keep'] for c in t['cells']]
    notes=[i['text'] for i in plan['items'] if i['kind'] in ('metadata','note') and i['text'] not in rendered_table_values]
    if notes:
        prototype=lanes[0]['caption_shape_id'];placements.append({'shape_id':prototype,'text':'\n'.join(notes),'col':0,'end_col':max(source.max_column,spans[-1][1]),'row':row+1,'end_row':row+8});row+=9
    target.print_area=f'A1:{get_column_letter(max(source.max_column,spans[-1][1]))}{row}'
    target.page_setup.orientation=source.page_setup.orientation;target.page_setup.paperSize=source.page_setup.paperSize;target.page_setup.fitToWidth=1;target.page_setup.fitToHeight=0;target.sheet_properties.pageSetUpPr.fitToPage=True
    target.sheet_view.showGridLines=False
    save(job.out/'native-textbox-placements.json',placements)
    return placements
