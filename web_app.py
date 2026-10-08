"""Local web interface for video guides and photo/description-to-Excel workflows."""
import argparse,json,mimetypes,os,re,subprocess,threading,time,uuid
from PIL import Image
import pymupdf
from excel_guide import discover,IMAGE_EXT
from openpyxl import load_workbook
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote,urlsplit

ROOT=Path(__file__).resolve().parent
OUTPUT=ROOT/'output'
EXTENSIONS={'.mp4','.mkv','.mov','.avi','.webm'}
LOCK=threading.Lock()
JOBS={}
PAGE=ROOT/'web_ui.html'

def videos():
    return [{'path':str(p.relative_to(ROOT)),'name':p.name} for base in [ROOT,OUTPUT/'web_uploads'] if base.exists() for p in sorted(base.iterdir()) if p.is_file() and p.suffix.lower() in EXTENSIONS]

def results():
    rows=[]
    if not OUTPUT.exists():return rows
    for folder in OUTPUT.iterdir():
        if not folder.is_dir() or folder.name=='web_uploads':continue
        files={ext:'/files/'+str((folder/f'guide-auto.{ext}').relative_to(OUTPUT)) for ext in ['html','pdf','md'] if (folder/f'guide-auto.{ext}').is_file()}
        if files:
            stamp=max((folder/f'guide-auto.{ext}').stat().st_mtime for ext in files)
            rows.append({'name':folder.name,'files':files,'modified':stamp})
    return sorted(rows,key=lambda r:r['modified'],reverse=True)

def launch(data):
    path=(ROOT/str(data.get('video',''))).resolve()
    permitted={str((ROOT/v['path']).resolve()) for v in videos()}
    if str(path) not in permitted:raise ValueError('Chọn một video trong danh sách hoặc tải video lên.')
    parallel=int(data.get('parallel',3));api_parallel=int(data.get('api_parallel',4))
    if not 1<=parallel<=8 or not 1<=api_parallel<=8:raise ValueError('Số luồng phải từ 1 đến 8.')
    language=str(data.get('language','')).strip()
    if language and not re.fullmatch('[a-z]{2,3}',language):raise ValueError('Mã ngôn ngữ không hợp lệ.')
    with LOCK:
        if busy():raise ValueError('Đang có tác vụ xử lý. Đợi tác vụ hiện tại kết thúc.')
        # Check CLI jobs as well: launching the web interface itself does not start GPU work.
        running=subprocess.run(['pgrep','-f','^'+str(ROOT/'.venv/bin/python')+' '+str(ROOT/'auto_guide')],capture_output=True,text=True)
        if running.returncode==0:raise ValueError('Pipeline đang chạy qua CLI. Đợi tác vụ đó kết thúc.')
        job_id=time.strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:8]
        out=OUTPUT/('web_'+job_id);out.mkdir(parents=True)
        command=[str(ROOT/'run_auto.sh'),'--video',str(path),'--output',str(out),'--parallel',str(parallel),'--api-parallel',str(api_parallel)]
        if language:command+=['--language',language]
        if data.get('second_pass'):command+=['--asr-second-pass']
        log_path=out/'web-run.log'
        with log_path.open('wb') as log:
            process=subprocess.Popen(command,cwd=ROOT,stdout=log,stderr=subprocess.STDOUT,start_new_session=True)
        JOBS[job_id]={'id':job_id,'kind':'video','video':path.name,'output':out,'process':process,'started':time.time()}
        remember(JOBS[job_id]);return job_id

def saved_job(job_id):
    # New jobs persist their metadata, so reloading the web server keeps progress accessible.
    for prefix in ['web_','excel_web_']:
        out=OUTPUT/(prefix+job_id);meta=out/'web-job.json'
        if meta.is_file():
            job=json.loads(meta.read_text());job['output']=out;return job
    # Recover videos started by the earlier interface, which did not persist metadata.
    out=OUTPUT/('web_'+job_id)
    if out.is_dir() and (out/'web-run.log').exists() and (out/'config.json').exists():
        config=json.loads((out/'config.json').read_text());pid=None
        found=subprocess.run(['pgrep','-f','^'+re.escape(str(ROOT/'.venv/bin/python'))+' '+re.escape(str(ROOT/'auto_guide.py'))],capture_output=True,text=True)
        for value in found.stdout.split():
            try:
                command=Path('/proc/'+value+'/cmdline').read_bytes().split(b'\0')
                if str(out).encode() in command:pid=int(value)
            except OSError:pass
        return {'id':job_id,'kind':'video','video':Path(config['video']).name,'output':out,'pid':pid,'started':(out/'config.json').stat().st_mtime}
    raise ValueError('Không tìm thấy tác vụ.')

def alive(job):
    if 'process' in job:return job['process'].poll() is None
    if not job.get('pid'):return False
    try:
        command=Path('/proc/'+str(job['pid'])+'/cmdline').read_bytes().split(b'\0')
        return str(job['output']).encode() in command
    except OSError:return False

def busy():
    if any(alive(j) for j in JOBS.values()):return True
    found=subprocess.run(['pgrep','-f','^'+re.escape(str(ROOT/'.venv/bin/python'))+' '+re.escape(str(ROOT))+r'/(auto_guide|excel_guide)\.py( |$)'],capture_output=True,text=True)
    return found.returncode==0

def remember(job):
    value={k:v for k,v in job.items() if k not in ('process','output')};value['pid']=job['process'].pid
    (job['output']/'web-job.json').write_text(json.dumps(value,ensure_ascii=False,indent=2))

def job_status(job_id):
    with LOCK:
        job=JOBS.get(job_id) or saved_job(job_id);out=job['output'];kind=job.get('kind','video');is_running=alive(job)
        artifact=out/('excel-result.json' if kind=='excel' else 'performance.json')
        result=json.loads(artifact.read_text()) if artifact.exists() else {}
        if kind=='excel' and not is_running and result.get('status')=='completed':code=0
        elif 'process' in job:code=job['process'].poll()
        else:code=None if is_running else 0 if result.get('status')=='completed' else 1
        elapsed=round(time.time()-job['started'])
        if not is_running:elapsed=round(result.get('elapsed_seconds',result.get('wall_seconds',elapsed)))
        info={'id':job_id,'kind':kind,'video':job.get('video',job.get('folder','')),'status':'running' if is_running else 'completed' if code==0 else 'failed',
              'elapsed_seconds':elapsed,'exit_code':code,'output':str(out.relative_to(ROOT)),'error':result.get('error')}
        if kind=='excel':
            info['files']={'xlsx':'/files/'+str((out/result['file']).relative_to(OUTPUT))} if result.get('file') and (out/result['file']).is_file() else {}
            info['summary']=result
            if (out/'matching-plan.json').exists():
                info['plan']=json.loads((out/'matching-plan.json').read_text())
                for item in info['plan']['items']:item['images']=[{'id':i,'url':'/files/'+str((out/'assets'/(i+'.jpg')).relative_to(OUTPUT))} for i in item.get('image_ids',[])]
        else:info['files']={ext:'/files/'+str((out/f'guide-auto.{ext}').relative_to(OUTPUT)) for ext in ['html','pdf','md'] if (out/f'guide-auto.{ext}').is_file()}
        log_path=out/('excel-run.log' if kind=='excel' else 'web-run.log')
        with log_path.open('rb') as f:
            f.seek(0,2);f.seek(max(0,f.tell()-30000));info['log']=f.read().decode('utf-8',errors='replace')
        return info

def inspect_excel(data):
    folder=Path(str(data.get('folder',''))).expanduser().resolve()
    if not folder.is_relative_to(ROOT):raise ValueError('Chọn thư mục dữ liệu trong repo video-to-guide.')
    folder,template,description,pdf,photos=discover(folder,**{k:data.get(k) or None for k in ['template','description','pdf','images']})
    wb=load_workbook(template,read_only=True);names=wb.sheetnames;wb.close()
    if len(names)<5:raise ValueError('Mẫu cần ít nhất 5 sheet.')
    return {'folder':str(folder),'template':str(template.relative_to(folder)),'description':str(description.relative_to(folder)),
            'pdf':str(pdf.relative_to(folder)),'photos':len(photos),'sheets':names,'description_text':description.read_text(encoding='utf-8-sig')}

def excel_datasets():
    rows=[]
    for base in [ROOT/'tests',OUTPUT/'excel_uploads']:
        if not base.exists():continue
        for folder in sorted(base.iterdir()):
            if not folder.is_dir():continue
            try:rows.append(inspect_excel({'folder':str(folder)}))
            except (ValueError,OSError):pass
    return rows

def excel_results():
    rows=[]
    for path in OUTPUT.glob('*/excel-result.json'):
        result=json.loads(path.read_text())
        if result.get('status')!='completed' or not result.get('file'):continue
        file=path.parent/result['file']
        if file.is_file():rows.append({'name':file.name,'url':'/files/'+str(file.relative_to(OUTPUT)),'modified':file.stat().st_mtime,'summary':result})
    return sorted(rows,key=lambda r:r['modified'],reverse=True)

def launch_excel(data):
    inspected=inspect_excel(data)
    with LOCK:
        if busy():raise ValueError('Đang có tác vụ xử lý. Đợi tác vụ hiện tại kết thúc.')
        job_id=time.strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:8];out=OUTPUT/('excel_web_'+job_id);out.mkdir(parents=True)
        command=[str(ROOT/'.venv/bin/python'),str(ROOT/'excel_guide.py'),'--folder',inspected['folder'],'--output',str(out)]
        for key in ['template','description','pdf','images']:
            if data.get(key):command+=['--'+key,str(data[key])]
        with (out/'excel-run.log').open('wb') as log:process=subprocess.Popen(command,cwd=ROOT,stdout=log,stderr=subprocess.STDOUT,start_new_session=True)
        job={'id':job_id,'kind':'excel','folder':inspected['folder'],'output':out,'process':process,'started':time.time()};JOBS[job_id]=job;remember(job)
        return job_id

UPLOAD_LOCK=threading.Lock()
UPLOAD_TYPES={'template':{'.xlsx'},'description':{'.txt'},'pdf':{'.pdf'},'image':IMAGE_EXT}

def create_excel_upload():
    token=uuid.uuid4().hex;folder=OUTPUT/'excel_uploads'/token;folder.mkdir(parents=True)
    (folder/'upload.json').write_text(json.dumps({'files':[],'ready':False}))
    return {'upload_id':token}

def upload_directory(token):
    if not re.fullmatch(r'[a-f0-9]{32}',token):raise ValueError('Mã upload không hợp lệ.')
    folder=OUTPUT/'excel_uploads'/token
    if not (folder/'upload.json').is_file():raise ValueError('Không tìm thấy bộ file upload.')
    return folder

def receive_excel_file(token,role,filename,stream,length):
    folder=upload_directory(token)
    if role not in UPLOAD_TYPES:raise ValueError('Loại file upload không hợp lệ.')
    if '/' in filename or '\\' in filename or filename in ('','.','..'):raise ValueError('Tên file không hợp lệ.')
    suffix=Path(filename).suffix.lower()
    if suffix not in UPLOAD_TYPES[role]:raise ValueError('Định dạng file không đúng loại đã chọn.')
    if not 0<length<=512*1024**2:raise ValueError('Mỗi file phải từ 1 byte đến 512 MB.')
    with UPLOAD_LOCK:
        manifest=json.loads((folder/'upload.json').read_text())
        if manifest['ready']:raise ValueError('Bộ file đã hoàn tất, hãy tạo lượt upload mới.')
        if sum(f['size'] for f in manifest['files'])+length>2*1024**3:raise ValueError('Bộ file vượt quá 2 GB.')
        if role!='image' and any(f['role']==role for f in manifest['files']):raise ValueError('Chỉ chọn một file cho mỗi loại mẫu/mô tả/PDF.')
        relative=Path('images')/(f'{sum(f["role"]=="image" for f in manifest["files"])+1:05d}_'+filename) if role=='image' else Path(role+suffix)
        target=folder/relative;target.parent.mkdir(exist_ok=True);temporary=target.with_name(target.name+'.upload')
        try:
            with temporary.open('wb') as handle:
                remaining=length
                while remaining:
                    block=stream.read(min(1024*1024,remaining))
                    if not block:raise ValueError('Upload chưa hoàn tất.')
                    handle.write(block);remaining-=len(block)
            if role=='image':
                with Image.open(temporary) as image:image.verify()
            elif role=='description':temporary.read_text(encoding='utf-8-sig')
            elif role=='template':
                # openpyxl validates the actual ZIP structure using a file handle.
                with temporary.open('rb') as handle:
                    workbook=load_workbook(handle,read_only=True)
                    if len(workbook.sheetnames)<5:raise ValueError('Mẫu cần ít nhất 5 sheet.')
                    workbook.close()
            elif role=='pdf':
                with pymupdf.open(stream=temporary.read_bytes(),filetype='pdf') as document:
                    if document.page_count<1 or document.needs_pass:raise ValueError('PDF không đọc được hoặc cần mật khẩu.')
            temporary.replace(target)
            manifest['files'].append({'role':role,'name':filename,'path':str(relative),'size':length})
            (folder/'upload.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
        except Exception as error:
            temporary.unlink(missing_ok=True)
            if not isinstance(error,ValueError):raise ValueError('Không đọc được nội dung file: '+filename) from error
            raise
    return {'name':filename,'uploaded_files':len(manifest['files'])}

def complete_excel_upload(token):
    folder=upload_directory(token)
    with UPLOAD_LOCK:
        manifest=json.loads((folder/'upload.json').read_text());roles={f['role'] for f in manifest['files']}
        if roles!=set(UPLOAD_TYPES):raise ValueError('Cần upload mẫu XLSX, mô tả TXT, PDF và ít nhất một ảnh.')
        inspected=inspect_excel({'folder':str(folder),'images':'images'})
        manifest['ready']=True;(folder/'upload.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
        return dict(inspected,upload_id=token,images='images')

class Handler(BaseHTTPRequestHandler):
    def reply(self,status,value):
        body=json.dumps(value,ensure_ascii=False).encode();self.send_response(status)
        self.send_header('Content-Type','application/json; charset=utf-8');self.send_header('Content-Length',str(len(body)));self.end_headers();self.wfile.write(body)

    def send_file(self,path):
        if not path.is_file():self.send_error(404);return
        self.send_response(200);self.send_header('Content-Type',mimetypes.guess_type(path.name)[0] or 'application/octet-stream')
        self.send_header('Content-Length',str(path.stat().st_size));self.end_headers()
        with path.open('rb') as f:
            while block:=f.read(1024*1024):self.wfile.write(block)

    def do_GET(self):
        route=urlsplit(self.path).path
        try:
            if route=='/':return self.send_file(PAGE)
            if route=='/api/videos':return self.reply(200,{'videos':videos()})
            if route=='/api/results':return self.reply(200,{'results':results()})
            if route=='/api/excel/datasets':return self.reply(200,{'datasets':excel_datasets()})
            if route=='/api/excel/results':return self.reply(200,{'results':excel_results()})
            if route.startswith('/api/jobs/'):return self.reply(200,job_status(route.rsplit('/',1)[-1]))
            if route.startswith('/files/'):
                path=(OUTPUT/unquote(route[len('/files/'):])).resolve()
                if not path.is_relative_to(OUTPUT.resolve()):return self.send_error(403)
                return self.send_file(path)
            self.send_error(404)
        except (ValueError,OSError) as error:self.reply(400,{'error':str(error)})

    def do_POST(self):
        route=urlsplit(self.path).path
        try:
            length=int(self.headers.get('Content-Length','0'))
            if route=='/api/excel/uploads':
                if length>16384:raise ValueError('Yêu cầu không hợp lệ.')
                if length:self.rfile.read(length)
                return self.reply(201,create_excel_upload())
            match=re.fullmatch(r'/api/excel/uploads/([a-f0-9]{32})(/complete)?',route)
            if match:
                if match[2]:
                    if length>16384:raise ValueError('Yêu cầu không hợp lệ.')
                    if length:self.rfile.read(length)
                    return self.reply(200,complete_excel_upload(match[1]))
                return self.reply(201,receive_excel_file(match[1],self.headers.get('X-Role',''),unquote(self.headers.get('X-Filename','')),self.rfile,length))
            if route=='/api/upload':
                filename=Path(unquote(self.headers.get('X-Filename','video.mp4'))).name
                suffix=Path(filename).suffix.lower()
                if suffix not in EXTENSIONS:raise ValueError('Định dạng video chưa được hỗ trợ.')
                if not 0<length<=10*1024**3:raise ValueError('Video phải có dung lượng từ 1 byte đến 10 GB.')
                directory=OUTPUT/'web_uploads';directory.mkdir(parents=True,exist_ok=True)
                path=directory/(uuid.uuid4().hex+suffix);temporary=path.with_suffix(suffix+'.upload')
                try:
                    with temporary.open('wb') as f:
                        remaining=length
                        while remaining:
                            block=self.rfile.read(min(1024*1024,remaining))
                            if not block:raise ValueError('Tải video chưa hoàn tất.')
                            f.write(block);remaining-=len(block)
                    temporary.replace(path)
                finally:temporary.unlink(missing_ok=True)
                return self.reply(201,{'path':str(path.relative_to(ROOT)),'name':filename})
            if route in ('/api/excel/inspect','/api/excel/jobs'):
                if not 0<length<=16384:raise ValueError('Yêu cầu không hợp lệ.')
                data=json.loads(self.rfile.read(length))
                if not isinstance(data,dict):raise ValueError('Yêu cầu không hợp lệ.')
                return self.reply(200,inspect_excel(data)) if route.endswith('/inspect') else self.reply(201,{'id':launch_excel(data)})
            if route=='/api/jobs':
                if not 0<length<=16384:raise ValueError('Yêu cầu không hợp lệ.')
                data=json.loads(self.rfile.read(length))
                if not isinstance(data,dict):raise ValueError('Yêu cầu không hợp lệ.')
                return self.reply(201,{'id':launch(data)})
            self.send_error(404)
        except (ValueError,OSError,TypeError) as error:self.reply(400,{'error':str(error)})

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--host',default='127.0.0.1');parser.add_argument('--port',type=int,default=7860)
    args=parser.parse_args();server=ThreadingHTTPServer((args.host,args.port),Handler)
    print(f'Video-to-guide web: http://{args.host}:{args.port}',flush=True)
    server.serve_forever()
