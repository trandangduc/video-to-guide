import tempfile,unittest,zipfile
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font
from openpyxl.drawing.image import Image as XLImage
from PIL import Image
from excel_guide import source_items,transplant,ExcelJob

class ExcelPreservation(unittest.TestCase):
    def test_description_order_numbers_and_wrapped_lines(self):
        text='ABC123\n\n- Sử dụng 0.25m.\n  Không đổi đơn vị.\n\n- Kết quả cuối.\n'
        items=source_items(text)
        self.assertEqual([i['text'] for i in items],['ABC123','Sử dụng 0.25m. Không đổi đơn vị.','Kết quả cuối.'])
        self.assertEqual([i['id'] for i in items],['TXT_001','TXT_002','TXT_003'])

    def test_result_sheet_without_photos_is_rejected(self):
        job=ExcelJob.__new__(ExcelJob)
        with self.assertRaisesRegex(ValueError,'Sheet kết quả phải có ảnh'):
            job.workbook({'items':[{'kind':'result','image_ids':[]}]},[],[])

    def test_qwen38_pairs_one_photo_to_every_operation(self):
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary);(root/'assets').mkdir();images={}
            for index in range(2):
                path=root/f'photo-{index}.png';Image.new('RGB',(90,60),(index*80,40,100)).save(path)
                ident=f'IMG_{index+1:03d}';images[ident]={'id':ident,'path':str(path)}
            job=ExcelJob.__new__(ExcelJob);job.out=root;job.image_map=images
            plan={'items':[{'source_id':'TXT_001','kind':'step','text':'First action','image_ids':[]},
                           {'source_id':'TXT_002','kind':'step','text':'Second action','image_ids':[]} ]}
            job.ask=lambda *a,**k:{'pairs':[{'source_id':'TXT_001','image_id':'IMG_001','reason':'Visible state matches.','needs_review':False},
                                           {'source_id':'TXT_002','image_id':'IMG_002','reason':'Visible object matches.','needs_review':False}]}
            job.pair_every_step_photo(plan,[{'id':key,'description':'Observed photo'} for key in images])
            steps=[item for item in plan['items'] if item['kind']=='step']
            self.assertEqual([item['image_ids'] for item in steps],[['IMG_001'],['IMG_002']])
            self.assertTrue(all(item['checks'][0]['selected_by']=='qwen38-best-per-step' for item in steps))

    def test_transplant_preserves_original_sheets_images_formulas_and_hidden_sheet(self):
        with tempfile.TemporaryDirectory() as root:
            root=Path(root);image=root/'sample.png';Image.new('RGB',(90,60),'red').save(image)
            original=Workbook();original.remove(original.active)
            for i in range(6):
                ws=original.create_sheet(f'Page {i+1}');ws['A1']=f'original {i}';ws['B2']='=1+2';ws['A1'].font=Font(name='Arial',bold=True,color='FF0000');ws.add_image(XLImage(str(image)),'C3');ws.print_area='A1:H20'
            original.worksheets[5].sheet_state='hidden';original.save(root/'template.xlsx')
            generated=Workbook();generated.remove(generated.active)
            for i in range(3):
                ws=generated.create_sheet(f'Page {i+3}');ws['A1']=f'new {i}';ws['A1'].font=Font(name='Calibri',color='0000FF');ws.add_image(XLImage(str(image)),'E5');ws.print_area='A1:L100';ws.print_title_rows='1:8'
            generated.save(root/'new.xlsx')
            result=transplant(root/'template.xlsx',root/'new.xlsx',root/'result.xlsx')
            self.assertTrue(result['sheet_1_2_byte_identical']);self.assertTrue(result['original_styles_unchanged'])
            with zipfile.ZipFile(root/'template.xlsx') as a,zipfile.ZipFile(root/'result.xlsx') as b:
                for name in ['xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml','xl/drawings/drawing1.xml','xl/drawings/drawing2.xml','xl/media/image1.png','xl/worksheets/sheet6.xml']:
                    self.assertEqual(a.read(name),b.read(name),name)
            out=load_workbook(root/'result.xlsx');self.assertEqual(out.sheetnames,original.sheetnames)
            self.assertEqual(out.worksheets[0]['B2'].value,'=1+2');self.assertEqual(out.worksheets[0]['A1'].font.color.rgb,'00FF0000')
            self.assertEqual(out.worksheets[2]['A1'].value,'new 0');self.assertEqual(out.worksheets[2]['A1'].font.color.rgb,'000000FF')
            self.assertEqual(len(out.worksheets[2]._images),1);self.assertEqual(out.worksheets[5].sheet_state,'hidden')
            self.assertIn('$L$100',str(out.worksheets[2].print_area));self.assertIn('$H$20',str(out.worksheets[0].print_area))

if __name__=='__main__':unittest.main()

class HeaderPreservation(unittest.TestCase):
    def test_original_header_cells_merges_widths_and_logo_survive_body_replacement(self):
        from lxml import etree
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary);photo=root/'logo.png';Image.new('RGB',(184,72),'gray').save(photo)
            template=Workbook();template.remove(template.active)
            for i in range(5):
                ws=template.create_sheet('WI_'+str(i+1));ws.merge_cells('C1:H4');ws['C1']='MY HEADER';ws['J1']='ORIGINAL PN';ws['C1'].font=Font(name='Times New Roman',bold=True,size=22)
                ws.column_dimensions['C'].width=12.75;ws.row_dimensions[1].height=18;ws.add_image(XLImage(str(photo)),'A1');ws['B9']='old body'
            template.save(root/'template.xlsx')
            generated=Workbook();generated.remove(generated.active)
            for i in range(3):
                ws=generated.create_sheet('WI_'+str(i+3));ws['C1']='wrong replacement';ws['J1']='wrong code';ws['B9']='new body';ws.add_image(XLImage(str(photo)),'G9')
            generated.save(root/'generated.xlsx')
            transplant(root/'template.xlsx',root/'generated.xlsx',root/'result.xlsx',preserve_header=True)
            result=load_workbook(root/'result.xlsx')
            for ws in result.worksheets[2:5]:
                self.assertEqual(ws['C1'].value,'MY HEADER');self.assertEqual(ws['J1'].value,'ORIGINAL PN');self.assertEqual(ws['C1'].font.size,22)
                self.assertEqual(ws.column_dimensions['C'].width,12.75);self.assertEqual(ws.row_dimensions[1].height,18)
                self.assertIn('C1:H4',[str(r) for r in ws.merged_cells.ranges]);self.assertEqual(ws['B9'].value,'new body')
                self.assertEqual(len(ws._images),2)
                logo=next(i for i in ws._images if i.anchor._from.row==0)
                self.assertEqual(logo.width,184);self.assertEqual(logo.height,72);self.assertEqual(logo._data(),photo.read_bytes())
            ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            with zipfile.ZipFile(root/'template.xlsx') as a,zipfile.ZipFile(root/'result.xlsx') as b:
                for i in range(3,6):
                    name=f'xl/worksheets/sheet{i}.xml'
                    old=etree.fromstring(a.read(name));new=etree.fromstring(b.read(name))
                    for r in range(1,5):
                        before=old.find(f's:sheetData/s:row[@r="{r}"]',ns);after=new.find(f's:sheetData/s:row[@r="{r}"]',ns)
                        if before is not None:self.assertEqual(etree.tostring(before,method='c14n'),etree.tostring(after,method='c14n'))
